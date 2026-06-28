#!/usr/bin/env python3
"""Génère le livrable Excel des écritures comptables, avec contrôles bloquants.

Usage : python3 ecrire_livrable.py <ecritures.json> <livrable.xlsx>

Format du JSON d'entrée :
{
  "totaux_releve": {"sorties": 1234.56, "entrees": 789.00},   # totaux du relevé d'origine
  "ecritures": [
    {
      "piece": "2026-05-001-A", "journal": "AC", "date": "27/05/2026",
      "libelle": "ORISHA PM - Fact. 1400FC26035252",
      "lignes": [
        {"compte": "627000", "intitule": "Logiciels", "auxiliaire": "", "debit": 87.05, "credit": 0},
        {"compte": "445660", "intitule": "TVA déductible", "auxiliaire": "", "debit": 17.41, "credit": 0},
        {"compte": "419401", "intitule": "Fournisseurs gérance", "auxiliaire": "ORISHAPM", "debit": 0, "credit": 104.46}
      ]
    }
  ],
  "comptes_a_creer": [{"compte": "421000", "intitule": "Personnel - rémunérations dues", "justification": "..."}],
  "points_a_verifier": [{"piece": "...", "date": "...", "montant": 0.0, "sujet": "...", "detail": "..."}]
}

Le script REFUSE de produire le fichier si une pièce est déséquilibrée ou si le
bouclage banque (mouvements 512000 vs totaux du relevé) ne tombe pas juste.
"""
import json
import sys
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def D(x):
    return Decimal(str(x or 0)).quantize(Decimal("0.01"))


def style_header(ws, ncols):
    for j in range(1, ncols + 1):
        c = ws.cell(row=1, column=j)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def autosize(ws):
    for j, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(j)].width = min(width + 2, 60)


def main():
    if len(sys.argv) != 3:
        sys.exit("Usage: ecrire_livrable.py <ecritures.json> <livrable.xlsx>")
    with open(sys.argv[1], encoding="utf-8") as f:
        data = json.load(f)
    ecritures = data["ecritures"]

    # --- Contrôle 1 : équilibre de chaque pièce ---
    erreurs = []
    for e in ecritures:
        td = sum(D(l.get("debit")) for l in e["lignes"])
        tc = sum(D(l.get("credit")) for l in e["lignes"])
        if td != tc:
            erreurs.append(f"Pièce {e['piece']} déséquilibrée : D {td} / C {tc}")
    # --- Contrôle 2 : équilibre global ---
    tg_d = sum(D(l.get("debit")) for e in ecritures for l in e["lignes"])
    tg_c = sum(D(l.get("credit")) for e in ecritures for l in e["lignes"])
    if tg_d != tg_c:
        erreurs.append(f"Déséquilibre global : D {tg_d} / C {tg_c}")
    # --- Contrôle 3 : bouclage banque 512000 vs relevé ---
    bq_d = sum(D(l.get("debit")) for e in ecritures for l in e["lignes"] if str(l["compte"]).startswith("512"))
    bq_c = sum(D(l.get("credit")) for e in ecritures for l in e["lignes"] if str(l["compte"]).startswith("512"))
    tr = data.get("totaux_releve") or {}
    if tr:
        if D(tr.get("entrees")) != bq_d:
            erreurs.append(f"Bouclage banque : débits 512000 {bq_d} ≠ entrées du relevé {D(tr.get('entrees'))}")
        if D(tr.get("sorties")) != bq_c:
            erreurs.append(f"Bouclage banque : crédits 512000 {bq_c} ≠ sorties du relevé {D(tr.get('sorties'))}")
    if erreurs:
        print("CONTRÔLES EN ÉCHEC — fichier NON généré :")
        for err in erreurs:
            print("  -", err)
        sys.exit(1)

    wb = Workbook()

    # Onglet Écritures
    ws = wb.active
    ws.title = "Écritures"
    ws.append(["Pièce", "Journal", "Date", "Compte", "Auxiliaire", "Intitulé compte",
               "Libellé écriture", "Débit", "Crédit"])
    for e in ecritures:
        for l in e["lignes"]:
            ws.append([e["piece"], e["journal"], e["date"], str(l["compte"]),
                       l.get("auxiliaire", ""), l.get("intitule", ""),
                       l.get("libelle", e["libelle"]),
                       float(D(l.get("debit"))) or None,
                       float(D(l.get("credit"))) or None])
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=9):
        for c in row:
            c.number_format = "#,##0.00"
    style_header(ws, 9)
    autosize(ws)

    # Onglet Comptes à créer
    ws2 = wb.create_sheet("Comptes à créer")
    ws2.append(["Compte", "Intitulé", "Justification"])
    for c in data.get("comptes_a_creer", []):
        ws2.append([str(c["compte"]), c["intitule"], c.get("justification", "")])
    style_header(ws2, 3)
    autosize(ws2)

    # Onglet Points à vérifier
    ws3 = wb.create_sheet("Points à vérifier")
    ws3.append(["Pièce", "Date", "Montant", "Sujet", "Détail"])
    for p in data.get("points_a_verifier", []):
        ws3.append([p.get("piece", ""), p.get("date", ""), p.get("montant", ""),
                    p.get("sujet", ""), p.get("detail", "")])
        for c in ws3[ws3.max_row]:
            c.fill = WARN_FILL
    style_header(ws3, 5)
    autosize(ws3)

    # Onglet Contrôles
    ws4 = wb.create_sheet("Contrôles")
    ws4.append(["Contrôle", "Valeur", "Statut"])
    ws4.append(["Nombre de pièces", len(ecritures), "—"])
    ws4.append(["Total débits", float(tg_d), "OK"])
    ws4.append(["Total crédits", float(tg_c), "OK"])
    ws4.append(["Équilibre de chaque pièce", "vérifié au centime", "OK"])
    if tr:
        ws4.append(["Bouclage banque 512000 / relevé",
                    f"entrées {float(bq_d):.2f} / sorties {float(bq_c):.2f}", "OK"])
    style_header(ws4, 3)
    autosize(ws4)

    wb.save(sys.argv[2])
    print(f"Livrable généré : {sys.argv[2]}")
    print(f"  {len(ecritures)} pièces, D = C = {tg_d} €")
    print(f"  Comptes à créer : {len(data.get('comptes_a_creer', []))}, "
          f"points à vérifier : {len(data.get('points_a_verifier', []))}")


if __name__ == "__main__":
    main()
